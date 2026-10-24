"""Parse the mill's monthly attendance musters (JUL/AUG/SEP-26-ATT.xlsx).

Each sheet holds one or more blocks; every block has a date row, then a header
row (S.NO | E.NO | NAME | DEPT | AG | ...), then worker rows. Per day there is a
pair of columns: S (shift code) and OT (hours). After the days come the Excel's
own summary columns (D H F G X Y TOTAL, 7 TO 7, 8 H, ABSENT, LEFT, OT,
ALL TOTAL, WKD, <fraction>, 0T).
"""
import openpyxl, re, json, sys, warnings, datetime as dt
warnings.filterwarnings("ignore")

TOK = re.compile(r"^[A-Z]{1,3}\d{1,5}$")
SUMMARY = ["D", "H", "F", "G", "X", "Y", "TOTAL", "7 TO 7", "8 H", "ABSENT", "LEFT", "OT", "ALL TOTAL", "WKD"]


def s(v):
    return "" if v is None else str(v).strip()


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def parse_file(path, month):  # month = "2026-09"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=110, values_only=True))
        hdr_idx = [i for i, r in enumerate(rows) if any(s(v) == "E.NO" for v in r[:4])]
        for bi, h in enumerate(hdr_idx):
            hdr = [s(v) for v in rows[h]]
            daterow = rows[h - 1] if h > 0 else ()
            col = {}
            for i, v in enumerate(hdr):
                if v and v not in col:
                    col[v] = i
            # day columns: date cells in the date row belonging to `month`
            days = {}
            for i, v in enumerate(daterow):
                if isinstance(v, (dt.datetime, dt.date)) and v.strftime("%Y-%m") == month and v.day not in days:
                    days[v.day] = i  # S column; OT is i+1
            if not days:
                continue  # block for another month (copied sheet)
            last_day_col = max(days.values()) + 1
            summ = {}
            for i in range(last_day_col + 1, len(hdr)):
                v = hdr[i]
                if v in SUMMARY and v not in summ:
                    summ[v] = i
            if "WKD" in summ:
                summ["_frac"] = summ["WKD"] + 1
                summ["_oth"] = summ["WKD"] + 2
            end = hdr_idx[bi + 1] if bi + 1 < len(hdr_idx) else len(rows)
            eno = col.get("E.NO")
            for r in rows[h + 1:end]:
                tok = s(r[eno]) if eno is not None and eno < len(r) else ""
                if not TOK.match(tok):
                    continue
                g = lambda k: s(r[col[k]]) if k in col and col[k] < len(r) else ""
                rec = {
                    "month": month, "sheet": ws.title, "block": bi, "token": tok,
                    "name": g("NAME"), "dept": g("DEPT"), "ag": g("AG"), "shift": g("SHIFT"),
                    "page": g("PAGE"), "onroll": g("ONROLL"), "ondept": g("ON.DEPT") or g("ON-DEPT"),
                    "onag": g("ON.AG") or g("ON-AG"), "left": g("LEFT"),
                    "days": {}, "summary": {},
                }
                for d, c in sorted(days.items()):
                    code = s(r[c]) if c < len(r) else ""
                    ot = num(r[c + 1]) if c + 1 < len(r) else 0
                    if code or ot:
                        rec["days"][d] = [code, ot]
                for k, c in summ.items():
                    rec["summary"][k] = num(r[c]) if c < len(r) else 0
                out.append(rec)
    return out


if __name__ == "__main__":
    D = "C:/Users/aakas/Downloads/"
    files = {"2026-07": "JULY-26-ATT (1).xlsx", "2026-08": "AUG-26-ATT -.xlsx", "2026-09": "SEP-26-ATT.xlsx"}
    allrecs = []
    for m, f in files.items():
        recs = parse_file(D + f, m)
        print(m, f, "->", len(recs), "rows", file=sys.stderr)
        allrecs += recs
    json.dump(allrecs, open(sys.argv[1], "w"))
